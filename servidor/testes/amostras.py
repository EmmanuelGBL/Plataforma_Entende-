"""Arquivos de teste gerados por código.

Gerar em vez de versionar tem duas vantagens: o repositório é público e não
recebe material de terceiros (RN07), e cada amostra declara aqui mesmo o
defeito que existe para provocar.

O texto imita um material de Ciências do 5º ano, como o do protótipo.
"""

import io

from docx import Document
from fpdf import FPDF
from PIL import Image

CABECALHO = "Coleção Aprender Juntos - Ciências - 5º ano"

PARAGRAFO_1 = (
    "As plantas produzem o próprio alimento por meio de um processo chamado fotossíntese. "
    "Para isso, elas utilizam a luz do sol, a água retirada do solo pelas raízes e o gás "
    "carbônico presente no ar, que entra pelas folhas."
)
PARAGRAFO_2 = (
    "Durante a fotossíntese, a planta libera gás oxigênio no ambiente. Esse gás é "
    "fundamental para a respiração da maioria dos seres vivos, inclusive dos seres humanos."
)
ITENS = ["Luz do sol", "Água", "Gás carbônico"]


class _Pdf(FPDF):
    def header(self):
        self.set_font("Helvetica", "", 9)
        self.set_y(8)
        self.cell(0, 5, CABECALHO, align="C")
        self.set_y(25)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 9)
        self.cell(0, 5, str(self.page_no()), align="C")


def _corpo(pdf: FPDF, texto: str) -> None:
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, texto, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)


def pdf_didatico(paginas: int = 2) -> bytes:
    """Título (20 pt), subtítulo (14 pt), corpo (11 pt), lista, cabeçalho e
    número de página repetidos em todas as páginas."""
    pdf = _Pdf()
    for numero in range(1, paginas + 1):
        pdf.add_page()
        if numero == 1:
            pdf.set_font("Helvetica", "B", 20)
            pdf.multi_cell(0, 10, "Capítulo 3 - A fotossíntese", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 8, f"Parte {numero}: o que as plantas precisam", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        _corpo(pdf, PARAGRAFO_1)
        _corpo(pdf, PARAGRAFO_2)
        pdf.set_font("Helvetica", "", 11)
        for item in ITENS:
            # "-" e não "•": a fonte padrão do fpdf2 é latin-1, que não tem o marcador.
            pdf.multi_cell(0, 6, f"- {item}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
    return bytes(pdf.output())


def _imagem() -> Image.Image:
    return Image.new("RGB", (600, 800), (235, 235, 235))


def pdf_paginas_de_imagem(com_texto: list[bool]) -> bytes:
    """Uma página por item: True = página de texto, False = página só com
    imagem, como sai de um scanner sem OCR."""
    pdf = FPDF()
    for tem_texto in com_texto:
        pdf.add_page()
        if tem_texto:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, PARAGRAFO_1 + " " + PARAGRAFO_2, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.image(_imagem(), x=10, y=10, w=190)
    return bytes(pdf.output())


def pdf_com_paginas(total: int) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", "", 11)
    for numero in range(1, total + 1):
        pdf.add_page()
        pdf.multi_cell(0, 6, f"Atividade {numero}. {PARAGRAFO_1}", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def pdf_com_senha() -> bytes:
    pdf = FPDF()
    pdf.set_encryption(owner_password="dono", user_password="aluno")
    pdf.add_page()
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, PARAGRAFO_1, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def docx_didatico() -> bytes:
    documento = Document()
    documento.add_heading("Capítulo 3 - A fotossíntese", level=1)
    documento.add_heading("O que as plantas precisam", level=2)
    documento.add_paragraph(PARAGRAFO_1)
    for item in ITENS:
        documento.add_paragraph(item, style="List Bullet")
    documento.add_paragraph(PARAGRAFO_2)
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "Entra na planta"
    tabela.cell(0, 1).text = "Sai da planta"
    tabela.cell(1, 0).text = "Gás carbônico"
    tabela.cell(1, 1).text = "Gás oxigênio"
    return _salvar(documento)


def docx_quase_vazio() -> bytes:
    """Material de Word que é, na prática, uma imagem colada com legenda."""
    documento = Document()
    documento.add_paragraph("Figura 1.")
    return _salvar(documento)


def _salvar(documento) -> bytes:
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def png() -> bytes:
    saida = io.BytesIO()
    _imagem().save(saida, format="PNG")
    return saida.getvalue()
